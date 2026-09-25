from __future__ import annotations

import os
import tempfile
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

NAVY = "071B2E"

# ---------------------------------------------------------------- Arabic in PDF
# Amiri (SIL Open Font License, assets/fonts/Amiri-OFL.txt) is used for any text that contains Arabic.
import re as _re
import sys as _sys
from pathlib import Path as _Path
from xml.sax.saxutils import escape as _escape

_ARABIC = _re.compile(r"[\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF\uFB50-\uFDFF\uFE70-\uFEFF]")
_FONTS = {"ready": None}


def _font_folder():
    base = _Path(getattr(_sys, "_MEIPASS", _Path(__file__).resolve().parent))
    return base / "assets" / "fonts"


def arabic_fonts():
    """Register Amiri once; returns (regular, bold) or (None, None) when the font files are missing."""
    if _FONTS["ready"] is None:
        try:
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            folder = _font_folder()
            pdfmetrics.registerFont(TTFont("Amiri", str(folder / "Amiri-Regular.ttf")))
            pdfmetrics.registerFont(TTFont("Amiri-Bold", str(folder / "Amiri-Bold.ttf")))
            _FONTS["ready"] = ("Amiri", "Amiri-Bold")
        except Exception:
            _FONTS["ready"] = (None, None)
    return _FONTS["ready"]


def has_arabic(text):
    return bool(_ARABIC.search(str(text or "")))


def shape_arabic(text):
    """Join the Arabic letters and put the line in visual (right-to-left) order for the PDF."""
    text = str(text or "")
    if not has_arabic(text): return text
    try:
        import arabic_reshaper
        shaped = arabic_reshaper.reshape(text)
    except Exception:
        shaped = text
    try:
        from bidi.algorithm import get_display
        return get_display(shaped)
    except Exception:
        return " ".join(reversed(shaped.split(" ")))


def pdf_paragraph(text, style, bold=False):
    """A Paragraph that renders Arabic with the Arabic font, and anything else with the style's font."""
    regular, bold_font = arabic_fonts()
    if has_arabic(text) and regular:
        from reportlab.lib.styles import ParagraphStyle
        size = style.fontSize * 1.2
        arabic_style = ParagraphStyle(f"ar-{style.name}-{bold}", parent=style, fontName=bold_font if bold else regular, fontSize=size, leading=size * 1.4,
                                      alignment=2 if not _re.search(r"[A-Za-z]{3}", str(text)) else style.alignment)
        return Paragraph(_escape(shape_arabic(text)), arabic_style)
    return Paragraph(_escape(str(text or "")), style)

def export_excel(path, title, headers, rows):
    wb = Workbook(); ws = wb.active; ws.title = title[:31]
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    cell = ws.cell(1, 1, title); cell.font = Font(size=16, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=NAVY); cell.alignment = Alignment(horizontal="center")
    ws.cell(2, 1, f"Generated: {datetime.now():%d-%m-%Y %H:%M}")
    for c, header in enumerate(headers, 1):
        h = ws.cell(4, c, header); h.font = Font(bold=True, color="FFFFFF")
        h.fill = PatternFill("solid", fgColor=NAVY); h.alignment = Alignment(horizontal="center")
    for r, values in enumerate(rows, 5):
        for c, value in enumerate(values, 1): ws.cell(r, c, value)
    for index, column in enumerate(ws.columns, 1):
        width = min(35, max(12, max(len(str(c.value or "")) for c in column) + 2))
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A5"; ws.auto_filter.ref = f"A4:{ws.cell(4, len(headers)).coordinate}"
    wb.save(path)

def export_pdf(path, title, headers, rows):
    doc = SimpleDocTemplate(str(path), pagesize=landscape(A4), rightMargin=10*mm, leftMargin=10*mm, topMargin=10*mm, bottomMargin=10*mm)
    styles = getSampleStyleSheet()
    from reportlab.lib.styles import ParagraphStyle
    story = [pdf_paragraph(title, styles["Title"], True), Paragraph(f"Generated: {datetime.now():%d-%m-%Y %H:%M}", styles["Normal"]), Spacer(1, 6*mm)]
    cell = ParagraphStyle("plain-cell", parent=styles["Normal"], fontSize=7, leading=8.5)
    head = ParagraphStyle("plain-head", parent=cell, fontName="Helvetica-Bold", textColor=colors.white)
    data = [[pdf_paragraph(h, head, True) if has_arabic(h) else h for h in headers]] + \
           [[pdf_paragraph(value, cell) if has_arabic(value) else ("" if value is None else str(value)) for value in row] for row in rows]
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#071B2E")), ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"), ("FONTSIZE", (0,0), (-1,-1), 7),
        ("GRID", (0,0), (-1,-1), .25, colors.grey), ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#F3F6F8")]),
        ("ALIGN", (2,1), (-1,-1), "RIGHT"), ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("BOTTOMPADDING", (0,0), (-1,0), 7), ("TOPPADDING", (0,0), (-1,0), 7),
    ]))
    story.append(table); doc.build(story)

def print_rows(title, headers, rows):
    handle = tempfile.NamedTemporaryFile(prefix="SaberAccounting_", suffix=".pdf", delete=False)
    handle.close(); export_pdf(handle.name, title, headers, rows)
    if os.name != "nt": raise RuntimeError("Printing is available in the Windows application")
    os.startfile(handle.name, "print")
    return handle.name

def export_invoice_pdf(path, invoice, items, logo_path=None, company=None):
    """Invoice / debit note / credit note: Item, Description, Qty, Unit, Unit Price, Discount, Total; then Total, Discount,
    Total HT, VAT 11% (struck through for exports), Total, and the amount in words in English and Arabic."""
    from reportlab.lib.styles import ParagraphStyle
    from tafqeet import amount_in_words
    doc=SimpleDocTemplate(str(path),pagesize=A4,rightMargin=14*mm,leftMargin=14*mm,topMargin=12*mm,bottomMargin=12*mm)
    styles=getSampleStyleSheet(); story=[]; money=lambda v: f"{float(v or 0):,.2f}"
    if logo_path and os.path.exists(str(logo_path)): story.append(Image(str(logo_path),width=32*mm,height=32*mm))
    company=company or {}; company_name=company.get("company_name") or "SABER FOR AUDIT"
    company_line=" | ".join(value for value in (company.get("company_address"),company.get("company_phone"),company.get("company_email"),company.get("company_website")) if value)
    subtype=str(invoice.get("doc_subtype") or "invoice"); kind=str(invoice.get("kind") or "sale")
    title={"credit_note":"CREDIT NOTE | إشعار دائن","debit_note":"DEBIT NOTE | إشعار مدين"}.get(subtype,"TAX INVOICE | فاتورة ضريبية" if kind=="sale" else "PURCHASE INVOICE | فاتورة شراء")
    story.extend([pdf_paragraph(company_name if has_arabic(company_name) else company_name.upper(),styles["Title"],True)])
    if company_line: story.append(pdf_paragraph(company_line,styles["Normal"]))
    if company.get("company_mof"): story.append(Paragraph(f'MOF / VAT No.: {company["company_mof"]}',styles["Normal"]))
    story.extend([Spacer(1,5*mm),pdf_paragraph(title,styles["Heading1"],True)])
    info=[["Number",invoice["invoice_number"],"Date",invoice["invoice_date"]],["Customer" if kind=="sale" else "Supplier",invoice.get("party_name") or "","Due Date",invoice.get("due_date") or "-"],
          ["Currency",invoice["currency"],"Payment",str(invoice.get("payment_method") or invoice.get("payment_status") or "").title()]]
    cell=ParagraphStyle("inv-cell",parent=styles["Normal"],fontSize=8.5,leading=10.5)
    info_table=Table([[pdf_paragraph(str(v),cell,i % 2 == 0) for i,v in enumerate(row)] for row in info],colWidths=[26*mm,70*mm,24*mm,62*mm])
    info_table.setStyle(TableStyle([("GRID",(0,0),(-1,-1),.3,colors.HexColor("#c9cfd6")),("BACKGROUND",(0,0),(0,-1),colors.HexColor("#F3F6F8")),("BACKGROUND",(2,0),(2,-1),colors.HexColor("#F3F6F8"))]))
    story.extend([info_table,Spacer(1,5*mm)])
    headers=["Item","Description","Qty","Unit","Unit Price","Discount","Total"]
    rows=[]; gross_total=0.0
    for x in items:
        qty=float(x.get("quantity") or 0); price=float(x.get("unit_price") or 0); gross=float(x.get("gross_amount") or qty*price)
        percent=float(x.get("discount_percent") or 0); discount=float(x.get("discount_amount") or 0) or gross*percent/100
        net=gross-discount; gross_total+=net
        rows.append([x.get("item_code") or "",pdf_paragraph(x["description"],cell),f"{qty:g}",x.get("unit") or "",money(price),(f"{percent:g}%" if percent else money(discount) if discount else ""),money(net)])
    if not rows: rows=[["","Invoice total","1","",money(invoice["subtotal"]),"",money(invoice["subtotal"])]]; gross_total=float(invoice["subtotal"] or 0)
    table=Table([headers]+rows,repeatRows=1,colWidths=[22*mm,66*mm,14*mm,14*mm,23*mm,19*mm,24*mm])
    table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#071B2E")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),
        ("GRID",(0,0),(-1,-1),.4,colors.grey),("FONTSIZE",(0,0),(-1,-1),8.5),("ALIGN",(2,1),(-1,-1),"RIGHT"),("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F3F6F8")])]))
    story.extend([table,Spacer(1,4*mm)])
    discount_percent=float(invoice.get("invoice_discount_percent") or 0); discount_amount=float(invoice.get("invoice_discount_amount") or 0)
    if not discount_amount and discount_percent: discount_amount=gross_total*discount_percent/100
    subtotal=float(invoice.get("subtotal") or 0); vat=float(invoice.get("vat") or 0); total=float(invoice.get("total") or subtotal+vat)
    export=invoice.get("vat_treatment") in ("zero_rated","exempt")
    vat_label="VAT 11%" if not export else "<strike>VAT 11%</strike>&nbsp; Export - zero rated (Art. 19)" if invoice.get("vat_treatment")=="zero_rated" else "<strike>VAT 11%</strike>&nbsp; Exempt (Art. 16-17)"
    bold=ParagraphStyle("inv-total",parent=styles["Normal"],fontName="Helvetica-Bold",fontSize=10,alignment=2)
    plain=ParagraphStyle("inv-total-plain",parent=styles["Normal"],fontSize=9.5,alignment=2)
    totals=[[Paragraph("Total",plain),Paragraph(money(gross_total),plain)]]
    if discount_amount: totals.append([Paragraph(f"Discount{f' {discount_percent:g}%' if discount_percent else ''}",plain),Paragraph("- "+money(discount_amount),plain)])
    totals+= [[Paragraph("Total HT (before VAT)",bold),Paragraph(money(subtotal),bold)],[Paragraph(vat_label,plain),Paragraph(money(vat),plain)],
              [Paragraph(f"TOTAL {invoice['currency']}",bold),Paragraph(money(total),bold)]]
    totals_table=Table(totals,colWidths=[60*mm,32*mm],hAlign="RIGHT")
    totals_table.setStyle(TableStyle([("LINEABOVE",(0,-1),(-1,-1),1,colors.HexColor("#071B2E")),("BACKGROUND",(0,-1),(-1,-1),colors.HexColor("#E8EDF2")),("TOPPADDING",(0,0),(-1,-1),2)]))
    words=amount_in_words(total,invoice["currency"])
    story.extend([totals_table,Spacer(1,5*mm),Paragraph(f"<b>Amount in words:</b> {words['en']}",styles["Normal"]),pdf_paragraph(words["ar"],styles["Normal"]),Spacer(1,4*mm)])
    if invoice.get("notes"): story.append(pdf_paragraph(f'Notes: {invoice["notes"]}',styles["Normal"]))
    if float(invoice.get("amount_paid") or 0): story.append(Paragraph(f'Amount paid: {money(invoice.get("amount_paid"))} &nbsp;&nbsp; Outstanding: {money(invoice.get("outstanding",total))}',styles["Normal"]))
    story.extend([Spacer(1,10*mm),Table([["Prepared by","Approved by","Received by"]],colWidths=[60*mm,60*mm,60*mm],style=[("LINEABOVE",(0,0),(-1,0),.5,colors.grey),("FONTSIZE",(0,0),(-1,-1),8),("ALIGN",(0,0),(-1,-1),"CENTER")])])
    doc.build(story)


# ---------------------------------------------------------------- multi-section official reports
def _plain(value):
    """Numbers stay numeric for Excel; Decimals become floats."""
    from decimal import Decimal
    if isinstance(value, Decimal): return float(value)
    return value


def _formatted(value):
    from decimal import Decimal
    if value is None: return ""
    if isinstance(value, bool): return str(value)
    if isinstance(value, int): return f"{value:,}"
    if isinstance(value, (float, Decimal)):
        number = float(value)
        return f"{number:,.0f}" if abs(number - round(number)) < 1e-9 else f"{number:,.2f}"
    return str(value)


def _safe_sheet_title(title):
    cleaned = "".join("-" if ch in '[]:*?/\\' else ch for ch in str(title))
    return cleaned[:31] or "Report"


def export_sections_excel(path, title, meta, sections):
    wb = Workbook(); ws = wb.active; ws.title = _safe_sheet_title(title)
    width = max([len(section["headers"]) for section in sections] + [2])
    navy = PatternFill("solid", fgColor=NAVY); total_fill = PatternFill("solid", fgColor="E8EDF2")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    cell = ws.cell(1, 1, title); cell.font = Font(size=15, bold=True, color="FFFFFF"); cell.fill = navy
    cell.alignment = Alignment(horizontal="center", vertical="center"); ws.row_dimensions[1].height = 26
    row = 2
    for line in list(meta or []) + [f"Generated: {datetime.now():%d-%m-%Y %H:%M}"]:
        ws.cell(row, 1, line).font = Font(italic=True, color="44546A"); row += 1
    widths = {}
    for section in sections:
        row += 1
        ws.cell(row, 1, section["heading"]).font = Font(size=12, bold=True, color=NAVY); row += 1
        for column, header in enumerate(section["headers"], 1):
            h = ws.cell(row, column, header); h.font = Font(bold=True, color="FFFFFF"); h.fill = navy
            h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            widths[column] = max(widths.get(column, 10), min(22, len(str(header)) + 2))
        ws.row_dimensions[row].height = 30; row += 1
        totals = set(section.get("total_rows") or [])
        for index, values in enumerate(section["rows"]):
            for column, value in enumerate(values, 1):
                c = ws.cell(row, column, _plain(value))
                if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                    c.number_format = "#,##0.00" if isinstance(c.value, float) and abs(c.value - round(c.value)) > 1e-9 else "#,##0"
                    c.alignment = Alignment(horizontal="right")
                if index in totals: c.font = Font(bold=True); c.fill = total_fill
                widths[column] = max(widths.get(column, 10), min(45, len(_formatted(value)) + 2))
            row += 1
    for column, value in widths.items(): ws.column_dimensions[get_column_letter(column)].width = value
    ws.sheet_view.showGridLines = True
    wb.save(path)


def export_sections_pdf(path, title, meta, sections):
    from reportlab.lib.styles import ParagraphStyle
    page = landscape(A4)
    doc = SimpleDocTemplate(str(path), pagesize=page, rightMargin=8*mm, leftMargin=8*mm, topMargin=10*mm, bottomMargin=12*mm, title=title)
    styles = getSampleStyleSheet()
    header_style = ParagraphStyle("header", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=6.5, leading=7.5, textColor=colors.white, alignment=1)
    cell_style = ParagraphStyle("cell", parent=styles["Normal"], fontSize=6.5, leading=7.5)
    story = [pdf_paragraph(title, styles["Title"], True)]
    for line in list(meta or []) + [f"Generated: {datetime.now():%d-%m-%Y %H:%M}"]: story.append(pdf_paragraph(line, styles["Normal"]))
    story.append(Spacer(1, 4*mm))
    available = page[0] - 16*mm
    for section in sections:
        story.append(pdf_paragraph(section["heading"], styles["Heading3"], True))
        headers = section["headers"]; count = len(headers)
        body = [[_formatted(value) for value in list(values) + [""] * (count - len(values))] for values in section["rows"]]
        weights = []
        def header_length(text):
            parts = [p.strip() for p in str(text).split(" | ")]
            return max(len(p) * (0.8 if has_arabic(p) else 1) for p in parts) * (0.75 if len(parts) > 1 else 1)
        for column in range(count):
            longest = max([header_length(headers[column])] + [len(row[column]) * (1.15 if has_arabic(row[column]) else 1) for row in body] or [6])
            weights.append(min(max(longest, 7), 30))
        scale = available / sum(weights); col_widths = [w * scale for w in weights]
        def header_cell(text):
            parts = [p.strip() for p in str(text).split(" | ")]
            if len(parts) == 1: return pdf_paragraph(parts[0], header_style, True)
            return [pdf_paragraph(p, header_style, True) for p in parts]  # English above, Arabic below
        data = [[header_cell(h) for h in headers]]
        for row in body:
            data.append([pdf_paragraph(value, cell_style) if has_arabic(value) or (len(value) > 18 and not value.replace(",", "").replace(".", "").replace("-", "").isdigit()) else value for value in row])
        table = Table(data, repeatRows=1, colWidths=col_widths)
        style = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#071B2E")), ("FONTSIZE", (0, 1), (-1, -1), 6.5),
                 ("GRID", (0, 0), (-1, -1), .25, colors.grey), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                 ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F6F8")]),
                 ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]
        for column in range(count):
            if any(value and value.replace(",", "").replace(".", "").lstrip("-").isdigit() for value in (row[column] for row in body)):
                style.append(("ALIGN", (column, 1), (column, -1), "RIGHT"))
        for index in section.get("total_rows") or []:
            if 0 <= index < len(body):
                style += [("FONTNAME", (0, index + 1), (-1, index + 1), "Helvetica-Bold"), ("BACKGROUND", (0, index + 1), (-1, index + 1), colors.HexColor("#E8EDF2"))]
        table.setStyle(TableStyle(style)); story += [table, Spacer(1, 5*mm)]

    def footer(canvas, document):
        canvas.saveState(); canvas.setFont("Helvetica", 7); canvas.setFillColor(colors.HexColor("#5F6B76"))
        footer_title = title.split(" | ")[0] if has_arabic(title) else title
        canvas.drawString(8*mm, 6*mm, footer_title); canvas.drawRightString(page[0] - 8*mm, 6*mm, f"Page {document.page}")
        canvas.restoreState()
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
