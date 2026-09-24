"""Tests for version 1.12: payroll official reports, quarterly VAT, user expiry/permissions,
legal-document alerts, and a full standalone run including backup and restore."""
import socket
import sqlite3
import tempfile
import threading
import time
import unittest
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

from client import ApiClient
from database import Database
from payroll_reports import build_payroll_report, period_range
from report_export import export_sections_excel, export_sections_pdf
from server import run_server
import vat_return


def new_db(folder, name="test.db"):
    db = Database(Path(folder) / name); db.initialize("secret")
    user = db.user_for_token(db.login("admin", "secret")["token"])
    return db, user["id"]


def cell(section, row, header):
    return section["rows"][row][section["headers"].index(header)]


class PayrollOfficialReportsTest(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory(ignore_cleanup_errors=True); self.db, self.user = new_db(self.folder.name)
        db, user = self.db, self.user
        self.employee = db.save_employee({"employee_number": "1000", "full_name": "Rami Employee", "currency": "LBP", "base_salary": "60000000",
            "marital_status": "married", "children": 2, "mof_number": "MOF-1", "nssf_number": "NSSF-1"}, user)
        self.manager = db.save_employee({"employee_number": "2000", "full_name": "Maya Manager", "currency": "USD", "base_salary": "3000",
            "employee_group": "manager"}, user)
        first = db.payroll_settings_for("2025-01-01")
        first.update({"date_from": "01-01-2025", "employee_ceiling": "50000000", "medical_ceiling": "50000000", "family_ceiling": "20000000"})
        db.save_payroll_settings(first, user)
        second = dict(first); second.update({"date_from": "01-05-2025", "employee_ceiling": "140000000", "medical_ceiling": "140000000", "family_ceiling": "28000000"})
        db.save_payroll_settings(second, user)
        for month in range(1, 7):
            for employee in (self.employee, self.manager):
                extra = {}
                if month == 4 and employee is self.employee: extra = {"retro_salary": "6000000", "retro_from": "01-01-2025", "retro_to": "31-03-2025"}
                if month == 6: extra["thirteenth_month"] = "60000000" if employee is self.employee else "3000"
                if month == 3: extra.update({"transport": "2000000" if employee is self.employee else "50", "schooling": "1000000" if employee is self.employee else "0", "bonus": "0"})
                record = db.save_payroll({"employee_id": employee["id"], "period_date": f"28-{month:02d}-2025", **extra}, user)
                if month != 6 or employee is self.employee: db.post_payroll(record["id"], user)

    def tearDown(self): self.folder.cleanup()

    def test_effective_periods_chain_and_reject_overlap(self):
        periods = self.db.list_payroll_settings()
        self.assertEqual([(p["date_from"], p["date_to"]) for p in periods], [("2025-01-01", "2025-04-30"), ("2025-05-01", None)])
        middle = dict(periods[0]); middle.update({"date_from": "01-03-2025", "date_to": "31-05-2025"})
        with self.assertRaisesRegex(ValueError, "overlaps"): self.db.save_payroll_settings(middle, self.user)
        with self.assertRaisesRegex(ValueError, "decimal rate"): self.db.save_payroll_settings({**periods[1], "employee_nssf_rate": "3"}, self.user)

    def test_ceilings_follow_the_period_of_each_month(self):
        rows = {r["period_date"]: r for r in self.db.list_payroll() if r["employee_id"] == self.employee["id"]}
        self.assertEqual(float(rows["2025-04-28"]["employee_nssf"]), 1500000)  # 66M capped at 50M ceiling
        self.assertEqual(float(rows["2025-05-28"]["employee_nssf"]), 1800000)  # 60M under the new 140M ceiling
        self.assertGreater(float(rows["2025-04-28"]["retro_tax"]), 0)

    def test_r10_quarterly_separate_groups_nssf_and_rates(self):
        report = build_payroll_report(self.db, "R10", "quarterly", 2025, 2, "both")
        headings = [s["heading"] for s in report["sections"]]
        self.assertIn("Salary tax withheld - Employees", headings); self.assertIn("Salary tax withheld - Managers", headings)
        tax = report["sections"][headings.index("Salary tax withheld - Employees")]
        self.assertEqual(cell(tax, 0, "Months"), 3); self.assertEqual(cell(tax, 0, "Retro Salary"), 6000000)
        self.assertEqual(cell(tax, 0, "13th Salary"), 60000000); self.assertGreater(cell(tax, 0, "of which Retro Tax"), 0)
        nssf = report["sections"][headings.index("NSSF contributions - Employees")]
        self.assertEqual(cell(nssf, 0, "Employee NSSF (3%)"), 1500000 + 1800000 + 3600000)
        self.assertGreater(cell(nssf, 0, "Employer Medical"), 0); self.assertGreater(cell(nssf, 0, "Employer End-of-Service"), 0)
        managers = report["sections"][headings.index("Salary tax withheld - Managers")]
        self.assertEqual(cell(managers, 0, "Months"), 2)  # June manager payroll is still a draft
        with_drafts = build_payroll_report(self.db, "R10", "quarterly", 2025, 2, "manager", include_drafts=True)
        self.assertEqual(cell(with_drafts["sections"][0], 0, "Months"), 3)
        rates = report["sections"][-1]["rows"]
        self.assertEqual(rates[0][:2], ["01-01-2025", "30-04-2025"]); self.assertEqual(rates[1][:2], ["01-05-2025", "Open"])
        managers_usd = cell(managers, 0, "Gross")
        self.assertEqual(managers_usd, 2 * 3000 * 89500)  # converted to LBP

    def test_monthly_yearly_r5_r6_and_transport_schooling(self):
        march = build_payroll_report(self.db, "R10", "monthly", 2025, 3, "employee")
        self.assertEqual(cell(march["sections"][0], 0, "Transport"), 2000000); self.assertEqual(cell(march["sections"][0], 0, "Schooling"), 1000000)
        r5 = build_payroll_report(self.db, "R5", "yearly", 2025, 1, "employee")
        summary = r5["sections"][0]; items = {row[0]: row[1] for row in summary["rows"]}
        self.assertEqual(items["Number of employees"], 1); self.assertEqual(items["Retro Salary"], 6000000)
        r6 = build_payroll_report(self.db, "R6", "yearly", 2025, 1, "both")
        employee_sheet = next(s for s in r6["sections"] if "Rami Employee" in s["heading"])
        self.assertEqual(len(employee_sheet["rows"]), 7)  # six months + total
        self.assertEqual(employee_sheet["rows"][3][employee_sheet["headers"].index("Retro Period")], "01-01-2025 to 31-03-2025")
        with self.assertRaisesRegex(ValueError, "Quarter"): period_range("quarterly", 2025, 5)

    def test_exports_to_excel_and_pdf(self):
        report = build_payroll_report(self.db, "R6", "yearly", 2025, 1, "both")
        xlsx = Path(self.folder.name) / "r6.xlsx"; pdf = Path(self.folder.name) / "r6.pdf"
        export_sections_excel(xlsx, report["title"], report["meta"], report["sections"]); export_sections_pdf(pdf, report["title"], report["meta"], report["sections"])
        sheet = load_workbook(xlsx).active
        self.assertEqual(sheet["A1"].value, report["title"])
        self.assertTrue(pdf.read_bytes().startswith(b"%PDF")); self.assertGreater(pdf.stat().st_size, 3000)


class QuarterlyVatTest(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory(ignore_cleanup_errors=True); self.db, self.user = new_db(self.folder.name)
        db, user = self.db, self.user
        db.save_exchange_rate({"date_from": "01-01-2025", "date_to": "31-12-2025", "from_currency": "USD", "to_currency": "LBP", "rate": "89500"}, user)
        def invoice(date, party, kind, currency, amount, status="posted"):
            return db.create_manual_invoice({"invoice_date": date, "party_name": party, "kind": kind, "currency": currency, "status": status},
                                            [{"description": "Line", "quantity": 1, "unit_price": amount, "vat_rate": 11}], user)
        self.sale = invoice("15-02-2025", "Client", "sales", "USD", 1000)
        self.purchase = invoice("20-02-2025", "Supplier", "purchases", "USD", 300)
        self.car = invoice("21-02-2025", "Car Dealer", "purchases", "USD", 200)
        self.lbp_sale = invoice("01-03-2025", "Local Client", "sales", "LBP", 8950000)
        self.review = invoice("10-03-2025", "Draft Client", "sales", "USD", 500, "review")
        db.add_expense({"expense_date": "05-03-2025", "description": "Supplies", "currency": "USD", "with_vat_subtotal": "100", "vat": "11"}, user)
        db.add_expense({"expense_date": "06-03-2025", "description": "Entertainment", "currency": "USD", "with_vat_subtotal": "100", "vat": "11", "vat_recoverable": False}, user)
        party = db.save_party({"kind": "supplier", "name": "Overseas Supplier"}, user)
        case = db.save_document_case({"case_type": "customs", "document_date": "12-03-2025", "party_id": party["id"], "currency": "USD", "supplier_invoice_amount": "1000", "import_vat": "121"}, user)
        for role in ("supplier_invoice", "customs_declaration", "broker_invoice"): db.add_case_attachment(case["id"], role, f"{role}.pdf", "application/pdf", b"%PDF", user)
        db.post_document_case(case["id"], user)
        db.set_vat_recoverable("invoice", self.car, False, user)

    def tearDown(self): self.folder.cleanup()

    def test_q1_lines_categories_non_deductible_and_lbp(self):
        result = vat_return.build_vat_return(self.db, 2025, 1)
        usd = result["per_currency"]["USD"]
        self.assertEqual(float(usd["sales"]["vat"]), 110); self.assertEqual(float(usd["purchases"]["vat"]), 33)
        self.assertEqual(float(usd["expenses"]["vat"]), 11); self.assertEqual(float(usd["customs"]["vat"]), 121)
        self.assertEqual(float(usd["non_deductible"]["vat"]), 22 + 11); self.assertEqual(float(usd["net"]["vat"]), 110 - 165)
        self.assertEqual(float(result["per_currency"]["LBP"]["sales"]["vat"]), 984500)
        self.assertEqual(result["totals_lbp"]["net"], (110 - 165) * 89500 + 984500)
        self.assertEqual(result["review_excluded"], 1)
        self.assertEqual(vat_return.build_vat_return(self.db, 2025, 1, include_review=True)["per_currency"]["USD"]["sales"]["vat"], 165)
        self.assertEqual(vat_return.build_vat_return(self.db, 2025, 1, currency="LBP")["currency_filter"], "LBP")

    def test_non_deductible_reclass_keeps_ledger_equal_to_return(self):
        balance = {row["code"]: row["closing_balance"] for row in self.db.trial_balance()}
        result = vat_return.build_vat_return(self.db, 2025, 1)
        self.assertAlmostEqual(balance["442660000"], float(result["per_currency"]["USD"]["total_input"]["vat"]), places=2)
        self.assertAlmostEqual(sum(r["debit"] - r["credit"] for r in self.db.journal()), 0, places=2)
        self.db.set_vat_recoverable("invoice", self.car, True, self.user)
        self.assertEqual(float(vat_return.build_vat_return(self.db, 2025, 1)["per_currency"]["USD"]["purchases"]["vat"]), 55)
        with self.assertRaisesRegex(ValueError, "purchase and expense"): self.db.set_vat_recoverable("invoice", self.sale, False, self.user)

    def test_adjustments_lock_and_credit_carry_forward(self):
        vat_return.add_adjustment(self.db, {"year": 2025, "quarter": 1, "adjustment_type": "output", "currency": "USD", "amount": "10", "reason": "Late credit note"}, self.user, "admin")
        with self.assertRaisesRegex(ValueError, "reason"): vat_return.add_adjustment(self.db, {"year": 2025, "quarter": 1, "adjustment_type": "input", "currency": "USD", "amount": "5", "reason": ""}, self.user)
        saved = vat_return.save_return(self.db, 2025, 1, self.user, user_name="admin")
        self.assertEqual(saved["status"], "saved")
        self.assertEqual(saved["credit_carried_forward_lbp"], 45 * 89500 - 984500)
        with self.assertRaisesRegex(ValueError, "saved"): vat_return.add_adjustment(self.db, {"year": 2025, "quarter": 1, "adjustment_type": "output", "currency": "USD", "amount": "1", "reason": "Too late"}, self.user)
        self.db.create_manual_invoice({"invoice_date": "10-04-2025", "party_name": "Client", "kind": "sales", "currency": "USD", "status": "posted"},
                                      [{"description": "Q2", "quantity": 1, "unit_price": 1000, "vat_rate": 11}], self.user)
        q2 = vat_return.build_vat_return(self.db, 2025, 2)
        self.assertEqual(q2["credit_brought_forward_lbp"], saved["credit_carried_forward_lbp"])
        self.assertEqual(q2["payable_lbp"], 110 * 89500 - saved["credit_carried_forward_lbp"])
        self.db.create_manual_invoice({"invoice_date": "28-03-2025", "party_name": "Client", "kind": "sales", "currency": "USD", "status": "posted"},
                                      [{"description": "After filing", "quantity": 1, "unit_price": 10, "vat_rate": 11}], self.user)
        self.assertTrue(vat_return.build_vat_return(self.db, 2025, 1)["changed_since_saved"])
        vat_return.reopen_return(self.db, 2025, 1, self.user)
        self.assertEqual(vat_return.build_vat_return(self.db, 2025, 1)["status"], "not saved")

    def test_previous_fiscal_year_file_provides_q1_credit(self):
        other = tempfile.TemporaryDirectory(ignore_cleanup_errors=True); previous, user = new_db(other.name, "2024.db")
        previous.create_manual_invoice({"invoice_date": "10-11-2024", "party_name": "Supplier", "kind": "purchases", "currency": "LBP", "status": "posted"},
                                       [{"description": "Stock", "quantity": 1, "unit_price": 100000000, "vat_rate": 11}], user)
        vat_return.save_return(previous, 2024, 4, user)
        result = vat_return.build_vat_return(self.db, 2025, 1, previous_year_db=previous)
        self.assertEqual(result["credit_brought_forward_lbp"], 11000000); self.assertIn("Q4 2024", result["credit_source"])
        manual = vat_return.build_vat_return(self.db, 2025, 1, credit_brought_forward="5000000")
        self.assertEqual(manual["credit_brought_forward_lbp"], 5000000); other.cleanup()

    def test_vat_exports(self):
        title, meta, sections = vat_return.export_sections(vat_return.build_vat_return(self.db, 2025, 1))
        xlsx = Path(self.folder.name) / "vat.xlsx"; pdf = Path(self.folder.name) / "vat.pdf"
        export_sections_excel(xlsx, title, meta, sections); export_sections_pdf(pdf, title, meta, sections)
        values = [c.value for row in load_workbook(xlsx).active.iter_rows() for c in row if c.value]
        self.assertIn("Credit carried forward to next quarter", values)
        self.assertTrue(pdf.read_bytes().startswith(b"%PDF"))


class UsersAlertsAndRatesTest(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory(ignore_cleanup_errors=True); self.db, self.user = new_db(self.folder.name)

    def tearDown(self): self.folder.cleanup()

    def test_new_users_expire_after_one_year_and_can_be_renewed(self):
        user = self.db.save_user({"username": "clerk", "password": "clerk123", "role": "accountant"}, self.user)
        self.assertEqual(user["expires_at"], (datetime.now() + timedelta(days=365)).strftime("%Y-%m-%d"))
        self.assertIsNone(next(u for u in self.db.list_users() if u["username"] == "admin")["expires_at"])
        token = self.db.login("clerk", "clerk123")["token"]
        self.db.save_user({"id": user["id"], "username": "clerk", "role": "accountant", "expires_at": "01-01-2020"}, self.user)
        self.assertIsNone(self.db.user_for_token(token))
        with self.assertRaisesRegex(PermissionError, "expired on 01-01-2020"): self.db.login("clerk", "clerk123")
        renewed = self.db.save_user({"id": user["id"], "username": "clerk", "role": "accountant", "renew": True}, self.user)
        self.assertEqual(renewed["days_remaining"], 365); self.assertTrue(self.db.login("clerk", "clerk123"))

    def test_permissions_last_admin_and_validation(self):
        clerk = self.db.save_user({"username": "clerk", "password": "clerk123", "role": "accountant", "permissions": {"payroll": False, "vat": True}}, self.user)
        row = self.db.user_for_token(self.db.login("clerk", "clerk123")["token"])
        self.assertFalse(self.db.user_can(row, "payroll")); self.assertTrue(self.db.user_can(row, "vat"))
        with self.assertRaisesRegex(ValueError, "administrator must remain"):
            self.db.save_user({"id": 1, "username": "admin", "role": "viewer"}, self.user)
        with self.assertRaisesRegex(ValueError, "already used"): self.db.save_user({"username": "CLERK", "password": "another1", "role": "viewer"}, self.user)
        with self.assertRaisesRegex(ValueError, "6 characters"): self.db.save_user({"id": clerk["id"], "username": "clerk", "role": "viewer", "password": "123"}, self.user)

    def test_sessions_expire(self):
        token = self.db.login("admin", "secret")["token"]
        with self.db.connect() as db: db.execute("UPDATE sessions SET created_at=? WHERE token=?", ((datetime.now() - timedelta(hours=30)).astimezone().isoformat(), token))
        self.assertIsNone(self.db.user_for_token(token))

    def test_legal_document_alerts(self):
        party = self.db.save_party({"kind": "supplier", "name": "Supplier"}, self.user)
        today = datetime(2026, 9, 24)
        for kind, expiry in (("Contract", "01-09-2026"), ("MOF / VAT Certificate", "10-10-2026"), ("ID / Passport", "31-12-2027"), ("Other", "")):
            self.db.add_party_document(party["id"], {"document_type": kind, "expiry_date": expiry, "file_name": "f.pdf", "mime_type": "application/pdf"}, b"%PDF", self.user)
        alerts = self.db.legal_document_alerts(30, today.strftime("%d-%m-%Y"))
        self.assertEqual((alerts["expired"], alerts["expiring"]), (1, 1))
        self.assertEqual(alerts["items"][0]["document_type"], "Contract"); self.assertEqual(alerts["items"][0]["status"], "expired")

    def test_old_payroll_date_formats_are_migrated_and_displayed(self):
        employee = self.db.save_employee({"employee_number": "1000", "full_name": "Legacy", "currency": "LBP", "base_salary": "1000"}, self.user)
        with self.db.connect() as db:
            db.execute("UPDATE payroll_settings SET date_from='01012025'")
            db.execute("""INSERT INTO payroll_records(payroll_number,employee_id,period_date,currency,retro_from,created_at)
                VALUES('PAY-OLD-1',?,'30062025','LBP','01-01-2025','x'),('PAY-OLD-2',?,'June 2025','LBP',NULL,'x')""", (employee["id"], employee["id"]))
        self.db.initialize("secret")
        periods = {r["payroll_number"]: r for r in self.db.list_payroll()}
        self.assertEqual(periods["PAY-OLD-1"]["period_date"], "2025-06-30"); self.assertEqual(periods["PAY-OLD-1"]["retro_from"], "2025-01-01")
        self.assertEqual(periods["PAY-OLD-2"]["period_date"], "June 2025")  # unreadable text is kept, not lost
        self.assertEqual(self.db.list_payroll_settings()[0]["date_from"], "2025-01-01")
        from desktop import safe_display_date
        self.assertEqual(safe_display_date("2025-06-30"), "30-06-2025"); self.assertEqual(safe_display_date("June 2025"), "June 2025")

    def test_exchange_rate_lookup_is_chronological(self):
        self.db.save_exchange_rate({"date_from": "01-06-2025", "date_to": "01-06-2025", "from_currency": "AED", "to_currency": "LBP", "rate": "24000"}, self.user)
        self.db.save_exchange_rate({"date_from": "01-07-2025", "date_to": "01-07-2025", "from_currency": "AED", "to_currency": "LBP", "rate": "25000"}, self.user)
        self.assertEqual(self.db._converted_amount(1, "AED", "LBP", "30-06-2025"), 24000)
        self.assertEqual(self.db._converted_amount(1, "AED", "LBP", "2025-07-15"), 25000)


class StandaloneEndToEndTest(unittest.TestCase):
    """Runs the embedded data service exactly as the installed app does and drives it through the API."""

    @classmethod
    def setUpClass(cls):
        cls.folder = tempfile.TemporaryDirectory(ignore_cleanup_errors=True); cls.database = Path(cls.folder.name) / "SaberAccounting" / "saber.db"
        cls.database.parent.mkdir(parents=True)
        probe = socket.socket(); probe.bind(("127.0.0.1", 0)); cls.port = probe.getsockname()[1]; probe.close()
        threading.Thread(target=run_server, kwargs={"host": "127.0.0.1", "port": cls.port, "database": str(cls.database), "admin_password": "admin123"}, daemon=True).start()
        cls.url = f"http://127.0.0.1:{cls.port}"
        for _ in range(100):
            try: ApiClient(cls.url).login("admin", "admin123"); break
            except Exception: time.sleep(0.1)

    @classmethod
    def tearDownClass(cls): cls.folder.cleanup()

    def admin(self):
        api = ApiClient(self.url); api.login("admin", "admin123")
        company = api.companies()[0]; api.select_company_year(company["id"], company["years"][0]["year"]); return api

    def test_full_workflow_permissions_backup_and_restore(self):
        api = self.admin(); year = api.fiscal_year
        employee = api.save_employee({"employee_number": "3000", "full_name": "E2E Employee", "currency": "LBP", "base_salary": "50000000"})
        for month in (1, 2, 3):
            api.post_payroll(api.save_payroll({"employee_id": employee["id"], "period_date": f"28-{month:02d}-{year}"})["id"])
        r10 = api.payroll_report("R10", "quarterly", year, 1, "employee")
        self.assertEqual(r10["record_count"], 3)
        api.create_manual_invoice({"invoice_date": f"10-02-{year}", "party_name": "E2E Client", "kind": "sales", "currency": "LBP", "status": "posted"},
                                  [{"description": "Service", "quantity": 1, "unit_price": 10000000, "vat_rate": 11}])
        self.assertEqual(api.vat_return(year, 1)["payable_lbp"], 1100000)
        api.save_user({"username": "nopay", "password": "nopay123", "role": "accountant", "permissions": {"payroll": False, "vat": True}})
        limited = ApiClient(self.url); session = limited.login("nopay", "nopay123"); limited.select_company_year(api.company_id, year)
        self.assertFalse(session["permissions"]["payroll"]); self.assertEqual(session["expires_at"], (datetime.now() + timedelta(days=365)).strftime("%Y-%m-%d"))
        with self.assertRaisesRegex(RuntimeError, "permission to use Payroll"): limited.payroll()
        self.assertEqual(limited.vat_return(year, 1)["payable_lbp"], 1100000)
        users = {u["username"]: u for u in api.users()}
        api.save_user({"id": users["nopay"]["id"], "username": "nopay", "role": "accountant", "expires_at": "01-01-2020"})
        with self.assertRaisesRegex(RuntimeError, "expired"): ApiClient(self.url).login("nopay", "nopay123")
        expired_calls = []; limited.on_unauthorized = lambda: expired_calls.append(True)
        with self.assertRaises(RuntimeError): limited.vat_returns()
        self.assertEqual(expired_calls, [True])
        backup = api.create_backup()["path"]
        api.create_manual_invoice({"invoice_date": f"11-02-{year}", "party_name": "After Backup", "kind": "sales", "currency": "LBP", "status": "posted"},
                                  [{"description": "Service", "quantity": 1, "unit_price": 20000000, "vat_rate": 11}])
        self.assertEqual(api.vat_return(year, 1)["payable_lbp"], 3300000)
        result = api.restore_backup(Path(backup).name)
        self.assertTrue(result["safety_backup"])
        self.assertEqual(api.vat_return(year, 1)["payable_lbp"], 1100000)
        self.assertEqual(api.payroll_report("R10", "quarterly", year, 1, "employee")["record_count"], 3)
        bad = Path(backup).parent / "saber_accounting_broken.db"; bad.write_bytes(b"not a database")
        with self.assertRaisesRegex(RuntimeError, "not a valid"): api.restore_backup(bad.name)
        connection = sqlite3.connect(str(self.database))
        try: self.assertEqual(connection.execute("PRAGMA integrity_check").fetchone()[0], "ok")
        finally: connection.close()


if __name__ == "__main__": unittest.main()
